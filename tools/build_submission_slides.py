"""Build a seven-slide presentation draft from the checked project records.

Requires python-pptx. The deck distinguishes completed checks from pending
research and does not substitute partial model checks for a 100-rollout result.
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def main():
    root = Path(__file__).resolve().parents[1]
    smoke = json.loads((root / "submission/model_smoke_check.json").read_text())
    slides = [
        ("Agibot Toy2Box", "BC 정책 재현과 평가 · 5분 발표 초안", [
            "고정 질문: Toy2Box에서 BC 정책은 왜 실패하는가?",
            "선택 모델: Mimic 499개 → BC-RNN/GMM",
            "학습 seed 101 · epoch 2000 · 최종 가중치 1개",
            "정식 평가·단계별 진단·통제 실험을 보완할 초안",
        ], "약 30초. 결과가 완성된 연구 발표가 아니라 제출 자료 초안임을 설명한다."),
        ("1. 문제 정의와 진단 기준", "과업 실패를 숫자로 설명한다", [
            "과업: 트럭을 집어 충분히 들고 박스 안에 배치",
            "실패 분류: 접근 / 파지 / 운반 / 배치",
            "성공률과 함께 어느 단계에서 실패했는지 측정",
            "다음 단계: baseline 100회 평가 후 주된 원인 하나 선택",
        ], "약 40초. 아직 완료하지 않은 실패 분류를 실제 결과처럼 설명하지 않는다."),
        ("2. 수집 데이터", "시범 수·길이·idle 비교", [
            "Dataset A: 기준 구성 100개 · idle 평균 0.70%",
            "Dataset C: 분산 멈춤 구성 20개 · idle 평균 19.79%",
            "수집 시범 통계와 학습 정책 평가 결과는 별도 지표",
        ], "약 40초. 데이터 방식과 개수, idle 정의를 설명한다. 이 비교만으로 인과 효과를 주장하지 않는다."),
        ("3. 학습 방법과 선택 모델", "학습 이력과 실제 입력을 확인했다", [
            "annotation 20개 → Mimic 성공 시범 499개",
            "BC-RNN: LSTM hidden 400, 2 layers, GMM 5 modes",
            "23차원 상태 입력 · 7차원 상대 EE 및 그리퍼 출력",
            "batch 100 · sequence length 10 · Adam lr 0.001",
            "사용자 지정 고정 epoch 2000 모델 1개와 SHA256 제공",
        ], "약 50초. 최고 성능 checkpoint를 검증해 고른 것이라고 주장하지 않는다."),
        ("4. 평가 방법과 확인 범위", "정식 평가와 동작 확인을 구분", [
            "검증 명령: seed 42 · rollout 100회 · horizon 800",
            "학습 seed 101·111의 성공률 평균 ± 표준편차 보고 필요",
            f"모델 동작 확인: 완료 {smoke['completed_rollouts']}회, 성공 {smoke['successes_completed']}회",
            "위 부분 실행은 100회 정식 성공률이 아니다",
            "선택한 시범 action 재생 6개는 기존 성공 조건 충족",
        ], "약 50초. 부분 실행 수치를 정식 평가나 두 학습 seed 비교 결과로 사용하지 않는다."),
        ("5. 통제 실험 설계 보완", "진단 → 가설 → 한 요인 변경 → 판정", [
            "100회 baseline에서 실패 단계와 상태를 먼저 분류",
            "원인 하나에 연결된 가설과 판정 기준을 실험 전에 고정",
            "시범 수·epoch·batch·checkpoint 선택 기준을 맞춤",
            "환경·초기 물체 분포·성공 판정·종료 조건 유지",
            "성공률 변화와 겨냥한 실패 유형 감소를 함께 확인",
        ], "약 45초. 연구 주제나 개선 결과가 확정된 것처럼 발표하지 않는다."),
        ("6. 결론과 제출 전 확인", "코드·모델·명령·발표 자료를 연결", [
            "준비: 환경·학습 코드, 가중치 1개, 평가 명령, 데이터 기록",
            "확인: 모델 이력·SHA256, 실제 시범 재생·영상 프레임 수",
            "보완: 학습 seed 111, 정식 평가, 실패 진단, 통제 실험",
            "PPT 결과·결론과 조원별 역할을 실제 실험 후 갱신",
            "평가 명령어·가중치 경로·실행 환경 확인",
        ], "약 45초. 실제 결과와 한계를 함께 정리하며 조원별 역할을 입력한다."),
    ]
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for index, (title, subtitle, bullets, notes) in enumerate(slides, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F4F7FB")

        def text(x, y, width, height, body, size, color="172B4D", bold=False):
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            frame = box.text_frame
            frame.word_wrap = True
            for line_index, line in enumerate(body.split("\n")):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name = "맑은 고딕"
                paragraph.font.size = Pt(size)
                paragraph.font.bold = bold
                paragraph.font.color.rgb = RGBColor.from_string(color)
                paragraph.space_after = Pt(14)
                paragraph.line_spacing = 1.15

        text(0.65, 0.45, 12, 0.65, title, 32, bold=True)
        text(0.68, 1.25, 12, 0.55, subtitle, 18, "4473A6")
        text(0.9, 2.05, 11.65, 4.7, "\n".join("• " + bullet for bullet in bullets), 22)
        text(0.7, 7.03, 11.8, 0.28, f"Robotics Simulation Assignment 2  |  발표 초안  |  {index}/7", 10, "67788D")
        slide.notes_slide.notes_text_frame.text = notes
    destination = root / "submission/toy2box_presentation_draft.pptx"
    presentation.save(destination)
    checked = Presentation(destination)
    assert len(checked.slides) == 7
    (root / "submission/presentation_outline.txt").write_text("\n\n".join(
        f"{title}\n{subtitle}\n" + "\n".join("- " + bullet for bullet in bullets) + "\n발표 메모: " + notes
        for title, subtitle, bullets, notes in slides) + "\n")
    print("Created:", destination)


if __name__ == "__main__":
    main()
